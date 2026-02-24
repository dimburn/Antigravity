from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# プレゼンテーションの作成
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# --- カラーパレット（より情緒的な色設計） ---
DEEP_BLACK = RGBColor(5, 5, 10)       # 漆黒（わずかに青み）
BLAZE_ORANGE = RGBColor(255, 60, 0)   # 炎のようなオレンジ
DAWN_BLUE = RGBColor(10, 20, 40)      # 夜明け前の群青
PLATINUM = RGBColor(230, 230, 230)    # 金属的な白
SUB_TEXT_GRAY = RGBColor(150, 150, 160) # 目立たないグレー

# ---------------------------------------------------------
# ヘルパー関数（シネマティック演出用）
# ---------------------------------------------------------

def add_cinematic_slide(prs, background_color=DEEP_BLACK):
    """映画のような黒背景スライドを作成"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = background_color
    return slide

def add_full_bg_placeholder(slide, text="[ここに情緒的な背景画像を配置]"):
    """全画面背景画像のプレースホルダー"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(30, 30, 30)
    shape.line.color.rgb = BLAZE_ORANGE
    shape.line.width = Pt(1)
    
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(14)
    p.font.color.rgb = SUB_TEXT_GRAY
    p.alignment = PP_ALIGN.CENTER
    return shape

def add_impact_text(slide, text, size=80, color=PLATINUM, align=PP_ALIGN.CENTER, top=Inches(2.5), bold=True):
    """中央にドンと置くインパクトテキスト"""
    box = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.33), Inches(4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = align

def add_cinematic_subtitle(slide, text, top=Inches(4), color=BLAZE_ORANGE):
    """映画の字幕のようなサブテキスト"""
    box = slide.shapes.add_textbox(Inches(2), top, Inches(9.33), Inches(1))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = f"- {text} -"
    p.font.size = Pt(24)
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    p.font.italic = True

def add_corner_label(slide, text, top=Inches(0.5), left=Inches(0.5)):
    """隅に小さく置くラベル（洗練された印象）"""
    box = slide.shapes.add_textbox(left, top, Inches(4), Inches(0.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text.upper()
    p.font.size = Pt(10)
    p.font.color.rgb = SUB_TEXT_GRAY
    p.font.bold = True
    p.space_after = 0

def add_note(slide, text):
    """発表者用ノート（演出指示）を追加"""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = f"【演出指示】\n{text}"

# ---------------------------------------------------------
# スライド生成 (全20枚：ストーリー構成)
# ---------------------------------------------------------

# --- Scene 1: Awakening (覚醒) ---

# S1: タイトルコール
s1 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s1, "【画像】暗闇に一本の光る道、または夜明け前の水平線\n（期待感、静寂）")
add_impact_text(s1, "From SUPPORT\nto INFRASTRUCTURE", size=72, top=Inches(2.5))
add_cinematic_subtitle(s1, "2026 KICKOFF PROJECT", top=Inches(5))
add_note(s1, "BGM：重低音の効いた静かなイントロ。山内氏、沈黙からゆっくり語り出す。")

# S2: 宣言（旧・組織改編）
s2 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s2, "【画像】エンジンが換装される瞬間、またはサナギから蝶へ\n（進化の瞬間）")
add_corner_label(s2, "Organization Strategy")
add_impact_text(s2, "REBORN", size=120, color=BLAZE_ORANGE, top=Inches(2))
add_cinematic_subtitle(s2, "食品流通から、リーテイルサポートへ。", top=Inches(4.5), color=PLATINUM)
add_note(s2, "「改編」ではなく「再生」。青山氏、力強く宣言。")

# S3: ターゲット（旧・目標）
s3 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s3, "【画像】荒波を越えていく巨大な船の船首、または嵐の中の灯台\n（困難と挑戦）")
add_impact_text(s3, "1,200", size=150, color=PLATINUM, top=Inches(1.5))
add_cinematic_subtitle(s3, "億という頂へ。", top=Inches(4.5), color=BLAZE_ORANGE)
# 小さく内訳
details = s3.shapes.add_textbox(Inches(4), Inches(5.5), Inches(5.33), Inches(2))
tf = details.text_frame
p = tf.paragraphs[0]
p.text = "SALES: 1200 / PROFIT: 23 / OP: 8.1"
p.font.size = Pt(14)
p.font.color.rgb = SUB_TEXT_GRAY
p.alignment = PP_ALIGN.CENTER
add_note(s3, "数字だけをドンと出す。細かい内訳は口頭で。「この数字は、我々のプライドです」")

# S4: 未来図 (Vision 2030)
s4 = add_cinematic_slide(prs, DAWN_BLUE) # 少し色を変える
add_full_bg_placeholder(s4, "【画像】サイバーパンクな未来都市の摩天楼\n（2030年の景色）")
add_corner_label(s4, "Vision 2030")
add_impact_text(s4, "2,000", size=140, align=PP_ALIGN.RIGHT, top=Inches(2))
# テキスト配置を工夫（右寄せの数字に対し、左側にメッセージ）
msg_box = s4.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(3))
tf = msg_box.text_frame
p = tf.paragraphs[0]
p.text = "想像できるか。\nグループの半翼を担う\n我々の未来を。"
p.font.size = Pt(36)
p.font.color.rgb = PLATINUM
p.font.bold = True
add_note(s4, "青山氏、問いかけるように。「遠い未来ではない、すぐそこの未来だ」")

# S5: 野望 (Vision 2040)
s5 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s5, "【画像】宇宙から見た地球の夜景（光のネットワーク）\n（スケール感）")
add_impact_text(s5, "1 TRILLION", size=90, color=BLAZE_ORANGE, top=Inches(2.5))
add_cinematic_subtitle(s5, "2040 Group Vision", color=PLATINUM, top=Inches(4.5))
add_note(s5, "山内氏、夢を語る。「1兆円企業の一員になる覚悟はあるか」")

# S6: ロードマップ
s6 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s6, "【画像】山頂へ続く険しい稜線と、そこを歩く登山隊\n（道のり）")
add_impact_text(s6, "THE PATH", size=60, align=PP_ALIGN.LEFT, top=Inches(1), color=BLAZE_ORANGE)
# グラフの代わりにマイルストーンを文字で配置
steps = ["2026: 1200", "2028: 1500", "2030: 2000"]
for i, step in enumerate(steps):
    box = s6.shapes.add_textbox(Inches(1 + i*4), Inches(3 + i*1), Inches(4), Inches(2))
    p = box.text_frame.paragraphs[0]
    p.text = step
    p.font.size = Pt(40 + i*10) # 徐々に大きく
    p.font.color.rgb = PLATINUM
    p.font.bold = True

# S7: 体制図（概念）
s7 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s7, "【画像】F1のピットクルー、またはSWATチームの集合写真シルエット\n（プロ集団）")
add_impact_text(s7, "4 UNITS", size=80, top=Inches(2))
add_cinematic_subtitle(s7, "DIV 1 / DIV 2 / SUPPORT / STRATEGY", color=PLATINUM, top=Inches(4))

# S8: キャリア
s8 = add_cinematic_slide(prs, DEEP_BLACK)
add_impact_text(s8, "YOUR STAGE", size=80, color=BLAZE_ORANGE, top=Inches(2))
add_cinematic_subtitle(s8, "ポストは用意した。あとは君が座るだけだ。", color=PLATINUM, top=Inches(4))

# S9: 中締め（ブリッジ）
s9 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s9, "【画像】扉が開いて光が差し込むシーン\n（次へ続く）")
add_impact_text(s9, "ARE YOU READY?", size=60)

# --- Scene 2: The Frontline (最前線) ---

# S10: 第1事業部 Intro
s10 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s10, "【画像】青山・佐藤・洲之内 3名のモノクロポートレート（映画ポスター風）")
add_corner_label(s10, "Division 1: Food Service")
add_impact_text(s10, "THE KING ROAD", size=60, align=PP_ALIGN.LEFT, top=Inches(5.5), color=BLAZE_ORANGE)
add_note(s10, "「王道を行く」")

# S11: 第1G
s11 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s11, "【画像】固い握手のクローズアップ（信頼）")
add_impact_text(s11, "PARTNER", size=90, top=Inches(2))
add_cinematic_subtitle(s11, "主要顧客シェア 100%へ", color=PLATINUM)

# S12: 第2G
s12 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s12, "【画像】地図の上に置かれたコンパスとペン（開拓）")
add_impact_text(s12, "FRONTIER", size=90, top=Inches(2))
add_cinematic_subtitle(s12, "未開のエリアを塗り替えろ", color=PLATINUM)

# S13: 第2事業部 Intro
s13 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s13, "【画像】山内・梶浦・平佐多 3名のポートレート（挑戦的な視線）")
add_corner_label(s13, "Division 2: Specialist")
add_impact_text(s13, "GAME CHANGER", size=60, align=PP_ALIGN.LEFT, top=Inches(5.5), color=BLAZE_ORANGE)

# S14: 第3G
s14 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s14, "【画像】ライトアップされたラグジュアリーホテルのロビー（気品）")
add_impact_text(s14, "QUALITY", size=90, top=Inches(2))
add_cinematic_subtitle(s14, "そのブランドに、相応しいインフラを", color=PLATINUM)

# S15: 第4G
s15 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s15, "【画像】温かい食卓の湯気、笑顔のシニア（温もり）")
add_impact_text(s15, "LIFE LINE", size=90, top=Inches(2))
add_cinematic_subtitle(s15, "日本の食卓を止めるな", color=PLATINUM)

# --- Scene 3: The Arsenal (武器庫) ---

# S16: 支援本部 Intro
s16 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s16, "【画像】巨大なサーバーラックや物流倉庫の整然とした美しさ（機能美）")
add_impact_text(s16, "THE ARSENAL", size=80, color=SUB_TEXT_GRAY, top=Inches(2))
add_cinematic_subtitle(s16, "最強の武器庫が、ここにある", color=PLATINUM)

# S17: 施策
s17 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s17, "【画像】世界地図と光るネットワーク線")
add_impact_text(s17, "6,000 NODES", size=90, top=Inches(1.5), color=BLAZE_ORANGE)
msg = s17.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11.33), Inches(2))
msg.text_frame.text = "そのネットワークは、\n誰にも真似できない。"
msg.text_frame.paragraphs[0].font.size = Pt(40)
msg.text_frame.paragraphs[0].font.color.rgb = PLATINUM
msg.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# S18: 戦略室 Intro
s18 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s18, "【画像】マトリックスのようなデジタルデータの奔流、またはチェス盤")
add_impact_text(s18, "THE BRAIN", size=80, color=SUB_TEXT_GRAY, top=Inches(2))
add_cinematic_subtitle(s18, "勝つための羅針盤", color=PLATINUM)

# S19: 戦略T
s19 = add_cinematic_slide(prs, DEEP_BLACK)
add_impact_text(s19, "DATA DRIVEN", size=70, top=Inches(2))
add_cinematic_subtitle(s19, "勘と経験を、科学する", color=PLATINUM)

# S20: グランドフィナーレ
s20 = add_cinematic_slide(prs, DEEP_BLACK)
add_full_bg_placeholder(s20, "【画像】太陽が昇る水平線、まばゆい光（夜明け）")
add_impact_text(s20, "CREATE\nTHE FUTURE", size=100, color=BLAZE_ORANGE, top=Inches(1.5))
add_cinematic_subtitle(s20, "さあ、共に創ろう。業界のインフラを。", color=PLATINUM, top=Inches(5.5))

# 保存
output_file = "2026_Kickoff_Cinematic.pptx"
prs.save(output_file)

print(f"Presentation saved to {output_file}")