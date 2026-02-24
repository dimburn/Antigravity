from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# プレゼンテーションの作成
prs = Presentation()

# スライドサイズを16:9に設定
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# カラー設定
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
ORANGE = RGBColor(255, 102, 0)  # ネオン感のあるオレンジ
GRAY = RGBColor(50, 50, 50)     # 画像プレースホルダー用
DARK_GRAY = RGBColor(30, 30, 30) # 補助的な背景色

# ---------------------------------------------------------
# ヘルパー関数定義
# ---------------------------------------------------------

def add_slide(prs):
    """空白のスライドを追加し、背景を黒にする"""
    slide_layout = prs.slide_layouts[6] # 空白レイアウト
    slide = prs.slides.add_slide(slide_layout)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BLACK
    return slide

def add_image_placeholder(slide, left, top, width, height, text="[ここに画像を挿入]"):
    """画像の置き場所を示すグレーのボックスを作成"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = GRAY
    shape.line.color.rgb = ORANGE
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    return shape

def add_title(slide, text, font_size=40, top=Inches(0.5), color=WHITE):
    """スライドタイトルの追加"""
    box = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.33), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT

def add_center_text(slide, text, font_size=60, color=WHITE, top=Inches(2.5), bold=True):
    """中央揃えのテキスト（スローガンなど）"""
    box = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.33), Inches(3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = PP_ALIGN.CENTER

def add_big_number(slide, number, label, left, top, width, height):
    """数値を強調表示するブロック"""
    # 数値
    num_box = slide.shapes.add_textbox(left, top, width, height)
    tf = num_box.text_frame
    p = tf.paragraphs[0]
    p.text = number
    p.font.size = Pt(80) # 特大
    p.font.color.rgb = ORANGE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # ラベル（単位など）
    label_box = slide.shapes.add_textbox(left, top + Inches(1.3), width, Inches(0.8))
    tf_l = label_box.text_frame
    p_l = tf_l.paragraphs[0]
    p_l.text = label
    p_l.font.size = Pt(20)
    p_l.font.color.rgb = WHITE
    p_l.alignment = PP_ALIGN.CENTER

def add_keyword_list(slide, items, left, top, width, height):
    """キーワードの箇条書き（アイコンの代わりに■を使用）"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    for item in items:
        p = tf.add_paragraph()
        p.text = f"■ {item}"
        p.font.size = Pt(28)
        p.font.color.rgb = WHITE
        p.space_after = Pt(20)

# ---------------------------------------------------------
# スライド生成メイン処理 (全21枚)
# ---------------------------------------------------------

# --- 導入・全体戦略パート (Slide 1-9) ---

# S1: オープニング
s1 = add_slide(prs)
add_image_placeholder(s1, Inches(0), Inches(0), Inches(13.33), Inches(7.5), "背景画像：光る橋・インフラ網")
# テキストを前面に配置
title_box = s1.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(2.5))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "From SUPPORT\nto INFRASTRUCTURE"
p.font.size = Pt(72)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

sub_box = s1.shapes.add_textbox(Inches(1), Inches(5), Inches(11.33), Inches(1))
tf_s = sub_box.text_frame
p_s = tf_s.paragraphs[0]
p_s.text = "～サポートから業界インフラへ変革～"
p_s.font.size = Pt(32)
p_s.font.color.rgb = ORANGE
p_s.font.bold = True
p_s.alignment = PP_ALIGN.CENTER

# S2: 組織改編
s2 = add_slide(prs)
add_title(s2, "組織改編：リーテイルサポート部門の新体制")
add_image_placeholder(s2, Inches(1), Inches(2.5), Inches(4.5), Inches(3.5), "旧：食品流通部門\n(グレー)")
add_image_placeholder(s2, Inches(7.8), Inches(2.5), Inches(4.5), Inches(3.5), "新：リーテイルサポート部門\n(オレンジ)")
# 矢印
arrow = s2.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6), Inches(4), Inches(1.3), Inches(0.8))
arrow.fill.solid()
arrow.fill.fore_color.rgb = WHITE

# S3: 2026目標
s3 = add_slide(prs)
add_title(s3, "2026年度 部門目標", color=ORANGE)
add_big_number(s3, "1,200", "売上高 (億円)", Inches(1), Inches(2.5), Inches(4), Inches(2))
add_big_number(s3, "23", "粗利 (億円)", Inches(5.5), Inches(2.5), Inches(3), Inches(2))
add_big_number(s3, "8.1", "営業利益 (億円)", Inches(9), Inches(2.5), Inches(3), Inches(2))

# S4: 2030ビジョン
s4 = add_slide(prs)
add_image_placeholder(s4, Inches(0), Inches(0), Inches(13.33), Inches(7.5), "背景画像：未来都市 2030")
add_title(s4, "2030 VISION", top=Inches(0.5))
add_center_text(s4, "売上高 2,000億円", font_size=88, top=Inches(2.5))
add_center_text(s4, "グループ全体4000億円の半翼を担う", font_size=32, top=Inches(4.5), bold=False)

# S5: 2040ビジョン
s5 = add_slide(prs)
add_image_placeholder(s5, Inches(0), Inches(0), Inches(13.33), Inches(7.5), "背景画像：宇宙・地球の夜明け")
add_title(s5, "2040 LONG-TERM VISION", top=Inches(0.5))
add_center_text(s5, "グループ 1兆円", font_size=100, color=ORANGE, top=Inches(2))
add_center_text(s5, "部門売上 5,000億円へ", font_size=48, top=Inches(4.5))

# S6: 5カ年計画
s6 = add_slide(prs)
add_title(s6, "5-Year Roadmap：着実な成長の軌跡")
add_image_placeholder(s6, Inches(1), Inches(2), Inches(11.33), Inches(4.5), "グラフ画像挿入\n(右肩上がりの折れ線：2026->2028->2030)")

# S7: 新組織図
s7 = add_slide(prs)
add_title(s7, "2026 新組織体制")
add_image_placeholder(s7, Inches(1), Inches(2), Inches(11.33), Inches(5), "組織図ツリー画像挿入\n(第1事業部 / 第2事業部 / 支援本部 / 戦略室)")

# S8: 人材体制
s8 = add_slide(prs)
add_title(s8, "人材拡充計画：2030年の組織イメージ")
add_image_placeholder(s8, Inches(1.5), Inches(2.5), Inches(4.5), Inches(4), "NOW 2026\n(小さいピラミッド)")
add_image_placeholder(s8, Inches(7.3), Inches(2), Inches(5), Inches(5), "FUTURE 2030\n(大きいピラミッド・ポスト増)")

# S9: スローガン再提示
s9 = add_slide(prs)
add_image_placeholder(s9, Inches(0), Inches(0), Inches(13.33), Inches(7.5), "背景画像：光の橋（明）")
add_center_text(s9, "From SUPPORT\nto INFRASTRUCTURE", font_size=60, color=ORANGE, top=Inches(2.5))

# --- 第1事業部パート (Slide 10-12) ---

# S10: 第1事業部表紙
s10 = add_slide(prs)
add_title(s10, "食品流通第1事業部", font_size=54, color=ORANGE)
add_title(s10, "外食産業No.1への布陣", font_size=32, top=Inches(1.5))
add_image_placeholder(s10, Inches(6), Inches(2), Inches(6.5), Inches(4.5), "写真：青山部門長・佐藤部長・洲之内統括")

# S11: 第1グループ
s11 = add_slide(prs)
add_title(s11, "第1グループ：主要顧客深耕 × 大型案件")
add_big_number(s11, "〇〇〇", "売上目標 (億円)", Inches(1), Inches(2), Inches(5), Inches(2))
add_keyword_list(s11, ["既存顧客シェア100%", "新規大型案件 奪取", "利益率改善 +Xpt"], Inches(1), Inches(4.5), Inches(11), Inches(3))
add_image_placeholder(s11, Inches(8), Inches(1.5), Inches(4.5), Inches(3), "イメージ：握手・旗")

# S12: 第2グループ
s12 = add_slide(prs)
add_title(s12, "第2グループ：差別化戦略 × エリア開拓")
add_big_number(s12, "〇〇〇", "売上目標 (億円)", Inches(1), Inches(2), Inches(5), Inches(2))
add_keyword_list(s12, ["地方有力チェーン開拓", "メニュー開発支援", "配送効率の極大化"], Inches(1), Inches(4.5), Inches(11), Inches(3))
add_image_placeholder(s12, Inches(8), Inches(1.5), Inches(4.5), Inches(3), "イメージ：地図・トラック")

# --- 第2事業部パート (Slide 13-15) ---

# S13: 第2事業部表紙
s13 = add_slide(prs)
add_title(s13, "食品流通第2事業部", font_size=54, color=ORANGE)
add_title(s13, "専門市場への挑戦", font_size=32, top=Inches(1.5))
add_image_placeholder(s13, Inches(6), Inches(2), Inches(6.5), Inches(4.5), "写真：山内副部門長・梶浦部長・平佐多統括")

# S14: 第3グループ
s14 = add_slide(prs)
add_title(s14, "第3グループ：ホテル・ブライダルの革新")
add_big_number(s14, "〇〇〇", "売上目標 (億円)", Inches(1), Inches(2), Inches(5), Inches(2))
add_keyword_list(s14, ["ハイエンドホテル攻略", "ブライダル新商品提案", "品質×コストの最適解"], Inches(1), Inches(4.5), Inches(11), Inches(3))
add_image_placeholder(s14, Inches(8), Inches(1.5), Inches(4.5), Inches(3), "イメージ：ホテル・指輪")

# S15: 第4グループ
s15 = add_slide(prs)
add_title(s15, "第4グループ：ヘルスケアフードのインフラ化")
add_big_number(s15, "〇〇〇", "売上目標 (億円)", Inches(1), Inches(2), Inches(5), Inches(2))
add_keyword_list(s15, ["施設向け提案営業 強化", "給食パートナー連携", "地域密着型ソリューション"], Inches(1), Inches(4.5), Inches(11), Inches(3))
add_image_placeholder(s15, Inches(8), Inches(1.5), Inches(4.5), Inches(3), "イメージ：介護施設・食事")

# --- 営業支援本部・経営戦略室パート (Slide 16-20) ---

# S16: 営業支援本部
s16 = add_slide(prs)
add_title(s16, "営業支援本部", font_size=54, color=ORANGE)
add_title(s16, "Professional Platform", font_size=32, top=Inches(1.5))
add_image_placeholder(s16, Inches(1), Inches(2.5), Inches(11.33), Inches(4.5), "組織図：3本の柱 (業務/購買/SCM)\n背景にギアやサーバー")

# S17: 支援本部施策
s17 = add_slide(prs)
add_title(s17, "2026 重点施策：インフラとしての『質』と『量』")
add_image_placeholder(s17, Inches(0), Inches(0), Inches(13.33), Inches(7.5), "背景：ネットワークマップ")
# テキストボックスを背景の上に
box17 = s17.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(5))
tf17 = box17.text_frame
items17 = [
    "【業務】No more 事務作業",
    "   新CRM・プロセス効率化",
    "",
    "【購買】圧倒的バイイングパワー",
    "   NW 6,000社超 × 原価低減",
    "",
    "【SCM】欠品ゼロ・遅延ゼロ",
    "   物流効率化 × 安定供給"
]
for item in items17:
    p = tf17.add_paragraph()
    p.text = item
    p.font.size = Pt(32)
    p.font.color.rgb = WHITE
    p.font.bold = True

# S18: 経営戦略室
s18 = add_slide(prs)
add_title(s18, "経営戦略室", font_size=54, color=ORANGE)
add_title(s18, "収益性 × 生産性 の最大化", font_size=32, top=Inches(1.5))
add_image_placeholder(s18, Inches(1), Inches(2.5), Inches(11.33), Inches(4.5), "イメージ：マトリックスコード / 分析グラフ")

# S19: 事業戦略チーム
s19 = add_slide(prs)
add_title(s19, "事業戦略チーム：Data & Knowledge")
add_keyword_list(s19, ["過去実績のフルデータベース化", "成功ナレッジの標準化・共有", "LTV最大化シミュレーション"], Inches(1), Inches(2.5), Inches(11), Inches(4))
add_image_placeholder(s19, Inches(8), Inches(4.5), Inches(4), Inches(2.5), "アイコン：DB・グラフ")

# S20: 人事戦略チーム
s20 = add_slide(prs)
add_title(s20, "人事戦略チーム：Human Capital")
add_keyword_list(s20, ["適材適所の戦略的配置", "採用力強化・リテンション", "メンタルサポート・エンゲージメント"], Inches(1), Inches(2.5), Inches(11), Inches(4))
add_image_placeholder(s20, Inches(8), Inches(4.5), Inches(4), Inches(2.5), "アイコン：人・ハート")

# --- グランドフィナーレ (Slide 21) ---

# S21
s21 = add_slide(prs)
add_image_placeholder(s21, Inches(0), Inches(0), Inches(13.33), Inches(7.5), "背景画像：夜明け・紙吹雪")
add_center_text(s21, "さあ、共に未来を創ろう！", font_size=64, top=Inches(2))
add_center_text(s21, "私たちの手で業界インフラを築こう！", font_size=40, color=ORANGE, top=Inches(4))

# 保存
output_file = "2026_Kickoff_Presentation.pptx"
prs.save(output_file)

print(f"Presentation saved to {output_file}")